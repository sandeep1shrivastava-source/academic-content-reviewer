# parser.py
import docx
import pandas as pd
from pypdf import PdfReader


def extract_document_text(uploaded_file) -> str:
    filename = uploaded_file.name.lower()

    if filename.endswith(".docx"):
        doc = docx.Document(uploaded_file)
        text_parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(
                    [cell.text.strip() for cell in row.cells]
                )
                if row_text.strip():
                    text_parts.append(row_text)
        return "\n".join(text_parts)

    elif filename.endswith((".xlsx", ".xls")):
        df = pd.read_excel(uploaded_file)
        return df.to_markdown(index=False)

    elif filename.endswith(".pdf"):
        reader = PdfReader(uploaded_file)
        return "\n".join(
            [
                page.extract_text()
                for page in reader.pages
                if page.extract_text()
            ]
        )

    elif filename.endswith(".txt"):
        return uploaded_file.read().decode("utf-8")

    raise ValueError(f"Unsupported file type: {filename}")