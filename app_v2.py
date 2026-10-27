import pandas as pd
import streamlit as st
import docx
from io import BytesIO
from parser import extract_document_text
from prompt_library import SOP_CONFIG

# Import from our NEW v2 engine
from reviewer_v2 import evaluate_content

st.set_page_config(page_title="Academic Content Reviewer V2", page_icon="🚀", layout="wide")

def create_docx_report(report_dict, score_val, dynamic_verdict):
    """Generates a downloadable Word Document from the JSON report data."""
    doc = docx.Document()
    doc.add_heading("Content Review & Quality Audit Report", 0)

    if report_dict.get("opening_statement"):
        doc.add_paragraph(report_dict.get("opening_statement"))

    doc.add_heading("1. Executive Summary", level=1)
    exec_sum = report_dict.get("executive_summary", {})
    doc.add_paragraph(f"Overall Compliance Score: {score_val}%")
    doc.add_paragraph(f"Critical Risks: {exec_sum.get('critical_risks', 'N/A')}")
    doc.add_paragraph(f"Structural Integrity: {exec_sum.get('structural_integrity', 'N/A')}")
    doc.add_paragraph(f"Factual Accuracy: {exec_sum.get('factual_accuracy', 'N/A')}")
    doc.add_paragraph(f"Primary Area for Improvement: {exec_sum.get('primary_area_for_improvement', 'N/A')}")

    doc.add_heading("2. Detailed Review Matrix", level=1)
    matrix = report_dict.get("detailed_review_matrix", [])
    if matrix:
        table = doc.add_table(rows=1, cols=4)
        table.style = 'Table Grid'
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = 'Criterion'
        hdr_cells[1].text = 'Score'
        hdr_cells[2].text = 'Status'
        hdr_cells[3].text = 'Observation'
        for item in matrix:
            row_cells = table.add_row().cells
            row_cells[0].text = str(item.get("criterion", ""))
            row_cells[1].text = str(item.get("score", ""))
            row_cells[2].text = str(item.get("status", ""))
            row_cells[3].text = str(item.get("observation", ""))

    doc.add_heading("3. Required Corrections", level=1)
    corrections = report_dict.get("required_corrections", [])
    if corrections:
        for i, corr in enumerate(corrections, 1):
            doc.add_paragraph(f"{i}. {corr.get('title', '')} [{corr.get('priority', '')}]", style='Heading 2')
            doc.add_paragraph(f"Location: {corr.get('location', '')}")
            doc.add_paragraph(f"Error: {corr.get('error', '')}")
            doc.add_paragraph(f"Correction: {corr.get('correction', '')}")
    else:
        doc.add_paragraph("No corrections required.")

    doc.add_heading("4. Final Verdict", level=1)
    doc.add_paragraph(dynamic_verdict)

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

st.title("🚀 Content Review System (Advanced V2)")

with st.sidebar:
    st.header("Settings & Credentials")
    api_key_input = st.text_input("Gemini API Key (Optional)", type="password")
    
    st.header("1. Core Inputs")
    category = st.selectbox("Content Category", options=list(SOP_CONFIG.keys()))
    subject = st.text_input("Course Title", placeholder="e.g., Environmental Science")
    uploaded_file = st.file_uploader("Upload Main Content (Required)", type=["docx", "xlsx", "xls", "pdf", "txt"])
    
    st.header("2. Advanced Comparison")
    uploaded_syllabus = st.file_uploader("Upload Syllabus (Optional)", type=["docx", "pdf", "txt"])
    custom_prompt = st.text_area("Custom AI Checker Rules (Optional)", placeholder="e.g., 'Check if all references are in APA 7th format.'")

    run_btn = st.button("Generate Advanced Review", type="primary", use_container_width=True)

if run_btn:
    if not subject.strip() or not uploaded_file:
        st.warning("Subject and Main Content Document are required.")
    else:
        with st.spinner("Executing advanced AI review..."):
            try:
                raw_text = extract_document_text(uploaded_file)
                syllabus_text = extract_document_text(uploaded_syllabus) if uploaded_syllabus else None
                key_to_pass = api_key_input.strip() if api_key_input.strip() else None
                
                # Call V2 evaluator
                report = evaluate_content(category, subject, raw_text, key_to_pass, syllabus_text, custom_prompt)

                st.divider()

                # Calculate Score strictly
                exec_sum = report.get("executive_summary", {})
                score = exec_sum.get("overall_compliance_score", report.get("final_score_percentage", 0))
                try:
                    score_val = float(str(score).replace("%", "").strip())
                except ValueError:
                    score_val = 0.0

                # Determine Strict Verdict (>95%)
                if score_val > 95:
                    dynamic_verdict = "VERDICT: Proceed for University Review (Score > 95%)"
                else:
                    dynamic_verdict = "VERDICT: Return to Content Development Team for Rework (Score <= 95%)"

                # Generate DOCX in memory
                docx_file = create_docx_report(report, score_val, dynamic_verdict)
                
                # --- RENDER UI ---
                col_title, col_btn = st.columns([3, 1])
                with col_title:
                    st.subheader("Executive Summary")
                with col_btn:
                    st.download_button(
                        label="📥 Download Report as Word (.docx)",
                        data=docx_file,
                        file_name=f"QA_Report_{subject.replace(' ', '_')}.docx",
                        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                    )

                col1, col2, col3 = st.columns([1, 1, 2])
                with col1:
                    st.metric("Overall Compliance", f"{score_val}%")
                with col2:
                    if score_val > 95:
                        st.success(dynamic_verdict)
                    else:
                        st.error(dynamic_verdict)
                with col3:
                    st.markdown(f"**Primary Focus:** {exec_sum.get('primary_area_for_improvement', 'N/A')}")

                st.divider()

                st.subheader("Detailed Review Matrix")
                matrix = report.get("detailed_review_matrix", [])
                if matrix:
                    df = pd.DataFrame(matrix)
                    col_map = {"criterion": "Criterion", "score": "Score (%)", "status": "Status", "observation": "Observation"}
                    df = df.rename(columns={k: v for k, v in col_map.items() if k in df})
                    st.dataframe(df, use_container_width=True, hide_index=True)

                st.divider()

                st.subheader("Required Corrections")
                corrections = report.get("required_corrections", [])
                if corrections:
                    for i, corr in enumerate(corrections, 1):
                        priority = corr.get("priority", "Medium")
                        p_tag = f"🔴 [{priority}]" if "High" in priority else f"🟡 [{priority}]"
                        st.markdown(f"#### {i}. {corr.get('title', 'Fix')} {p_tag}")
                        st.markdown(f"**Location:** `{corr.get('location', 'N/A')}`")
                        st.markdown(f"**Error:** {corr.get('error', 'N/A')}")
                        st.markdown(f"**Required Fix:** {corr.get('correction', 'N/A')}")
                        st.write("---")
                else:
                    st.success("No critical corrections required.")

            except Exception as e:
                st.error(f"Review failed: {str(e)}")
